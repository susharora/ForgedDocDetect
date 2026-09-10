#!/usr/bin/env python3
"""
Materialize the canonical FantasyID held-out-test manifest.

Purpose
-------
The official FantasyID test partition already exists in the frozen
Tech-1 discovery workbook. This tool converts ONLY that already-discovered
metadata into one small canonical CSV used later by every final test run.

Important scientific boundary
-----------------------------
This tool DOES NOT:

- open any test image;
- decode any test image;
- recompute any test image hash;
- open any test JSON;
- inspect test regions;
- load a model;
- load a checkpoint;
- perform inference;
- calculate metrics;
- derive or modify a threshold.

Therefore running this tool is NOT model access to held-out test data.

It only materializes metadata already frozen during Tech-1 discovery.

Order
-----
Rows retain their original order in the frozen workbook Images sheet
after filtering split == "test". No sorting is performed.

No print() is used.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import logging
import os
import re
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Mapping

import openpyxl
import yaml


# ======================================================================
# Repository
# ======================================================================

REPO_ROOT = Path(
    __file__
).resolve().parents[1]

if str(
    REPO_ROOT
) not in sys.path:

    sys.path.insert(
        0,
        str(
            REPO_ROOT
        ),
    )


LOGGER = logging.getLogger(
    "finalize_fantasyid_test_manifest"
)


SHA256_PATTERN = re.compile(
    r"^[0-9a-f]{64}$"
)


OUTPUT_COLUMNS = (
    "source_workbook_row",
    "image_path",
    "image_sha256",
    "file_stem",
    "traffic_type",
    "label",
    "variant",
    "hardware_source",
    "face_db",
    "face_id",
    "gender",
    "project_role",
)


REQUIRED_WORKBOOK_COLUMNS = {
    "split",
    "traffic_type",
    "variant",
    "hardware_source",
    "file_stem",
    "image_path",
    "image_sha256",
    "face_db",
    "face_id",
    "gender",
}


# ======================================================================
# Generic helpers
# ======================================================================

def resolve_repo_path(
    value: str | Path,
) -> Path:

    path = Path(
        value
    ).expanduser()

    if not path.is_absolute():

        path = (
            REPO_ROOT
            / path
        )

    return path.resolve()


def relative_repo_path(
    path: Path,
) -> str:

    return (
        path.resolve()
        .relative_to(
            REPO_ROOT
        )
        .as_posix()
    )


def sha256_file(
    path: Path,
) -> str:

    digest = hashlib.sha256()

    with path.open(
        "rb"
    ) as file:

        for chunk in iter(
            lambda: file.read(
                1024 * 1024
            ),
            b"",
        ):

            digest.update(
                chunk
            )

    return digest.hexdigest()


def load_yaml(
    path: Path,
) -> dict[str, Any]:

    if not path.is_file():

        raise FileNotFoundError(
            path
        )

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:

        value = (
            yaml.safe_load(
                file
            )
            or {}
        )

    if not isinstance(
        value,
        dict,
    ):

        raise TypeError(
            f"{path} must contain a YAML mapping."
        )

    return value


def require_mapping(
    value: Any,
    label: str,
) -> Mapping[Any, Any]:

    if not isinstance(
        value,
        Mapping,
    ):

        raise TypeError(
            f"{label} must be a mapping."
        )

    return value


def mapping_seed_value(
    mapping: Mapping[Any, Any],
    seed: int,
) -> Any:

    if seed in mapping:

        return mapping[
            seed
        ]

    text_seed = str(
        seed
    )

    if text_seed in mapping:

        return mapping[
            text_seed
        ]

    raise KeyError(
        f"Missing prerequisite for seed {seed}."
    )


def normalized_text(
    value: Any,
    *,
    label: str,
    allow_empty: bool = False,
) -> str:

    if value is None:

        text = ""

    else:

        text = str(
            value
        ).strip()

    if (
        not allow_empty
        and not text
    ):

        raise RuntimeError(
            f"Missing required workbook value: {label}"
        )

    return text


# ======================================================================
# Git gate
# ======================================================================

def require_clean_git() -> tuple[
    str,
    str,
]:

    commit = subprocess.run(
        [
            "git",
            "rev-parse",
            "HEAD",
        ],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()

    branch = subprocess.run(
        [
            "git",
            "rev-parse",
            "--abbrev-ref",
            "HEAD",
        ],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()

    status = subprocess.run(
        [
            "git",
            "status",
            "--porcelain",
            "--untracked-files=all",
        ],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()

    if status:

        raise RuntimeError(
            "Git working tree must be clean before the canonical "
            "held-out-test manifest is materialized.\n\n"
            f"{status}"
        )

    if len(
        commit
    ) != 40:

        raise RuntimeError(
            "Git HEAD is not a full 40-character SHA."
        )

    return (
        commit,
        branch,
    )


def require_git_tracked(
    path: Path,
) -> None:

    relative = relative_repo_path(
        path
    )

    result = subprocess.run(
        [
            "git",
            "ls-files",
            "--error-unmatch",
            relative,
        ],
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
    )

    if result.returncode != 0:

        raise RuntimeError(
            "Required prerequisite is not Git-tracked:\n"
            f"  {relative}"
        )


# ======================================================================
# Logging
# ======================================================================

def configure_logging(
    *,
    config: Mapping[Any, Any],
) -> Path:

    logging_cfg = require_mapping(
        config[
            "logging"
        ],
        "logging",
    )

    timestamp = (
        datetime.now(
            timezone.utc
        )
        .strftime(
            "%Y-%m-%d_%H%M%S_%fZ"
        )
    )

    directory = resolve_repo_path(
        logging_cfg[
            "directory"
        ]
    )

    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    path = (
        directory
        / str(
            logging_cfg[
                "filename"
            ]
        ).format(
            timestamp=timestamp
        )
    )

    LOGGER.handlers.clear()
    LOGGER.propagate = False

    level_name = str(
        logging_cfg[
            "level"
        ]
    ).upper()

    if not hasattr(
        logging,
        level_name,
    ):

        raise RuntimeError(
            f"Unknown logging level: {level_name}"
        )

    LOGGER.setLevel(
        getattr(
            logging,
            level_name,
        )
    )

    handler = logging.FileHandler(
        path,
        mode="x",
        encoding="utf-8",
    )

    formatter = logging.Formatter(
        "%(asctime)sZ | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%dT%H:%M:%S",
    )

    formatter.converter = (
        time.gmtime
    )

    handler.setFormatter(
        formatter
    )

    LOGGER.addHandler(
        handler
    )

    return path


# ======================================================================
# Frozen prerequisites
# ======================================================================

def validate_prerequisites(
    *,
    config: Mapping[Any, Any],
) -> dict[str, Any]:

    prerequisite_cfg = require_mapping(
        config[
            "prerequisites"
        ],
        "prerequisites",
    )

    protocol_cfg = require_mapping(
        prerequisite_cfg[
            "final_detection_protocol"
        ],
        "prerequisites.final_detection_protocol",
    )

    protocol_path = resolve_repo_path(
        protocol_cfg[
            "path"
        ]
    )

    require_git_tracked(
        protocol_path
    )

    expected_protocol_sha = str(
        protocol_cfg[
            "sha256"
        ]
    )

    actual_protocol_sha = sha256_file(
        protocol_path
    )

    if (
        actual_protocol_sha
        != expected_protocol_sha
    ):

        raise RuntimeError(
            "Final-detection protocol SHA mismatch:\n"
            f"  expected={expected_protocol_sha}\n"
            f"  actual={actual_protocol_sha}"
        )

    protocol = load_yaml(
        protocol_path
    )

    if (
        protocol.get(
            "status"
        )
        != "frozen_before_held_out_test_access"
    ):

        raise RuntimeError(
            "Final-detection protocol is not frozen."
        )

    selected = require_mapping(
        protocol[
            "selected_configuration"
        ],
        "final_detection.selected_configuration",
    )

    if selected[
        "resolution"
    ] != "r512":

        raise RuntimeError(
            "Frozen selected resolution must be r512."
        )

    if float(
        selected[
            "backbone_learning_rate"
        ]
    ) != 0.0001:

        raise RuntimeError(
            "Frozen backbone LR must equal 1e-4."
        )

    selected_seeds = tuple(
        int(
            value
        )
        for value
        in selected[
            "seeds"
        ]
    )

    if selected_seeds != (
        8,
        9,
        10,
    ):

        raise RuntimeError(
            "Frozen final seed set must be [8,9,10]."
        )

    thresholds_cfg = require_mapping(
        prerequisite_cfg[
            "dev_thresholds"
        ],
        "prerequisites.dev_thresholds",
    )

    threshold_evidence: dict[
        int,
        dict[str, Any],
    ] = {}

    for seed in (
        8,
        9,
        10,
    ):

        item = require_mapping(
            mapping_seed_value(
                thresholds_cfg,
                seed,
            ),
            f"prerequisites.dev_thresholds.{seed}",
        )

        path = resolve_repo_path(
            item[
                "path"
            ]
        )

        require_git_tracked(
            path
        )

        expected_sha = str(
            item[
                "sha256"
            ]
        )

        actual_sha = sha256_file(
            path
        )

        if actual_sha != expected_sha:

            raise RuntimeError(
                "Frozen dev-threshold SHA mismatch:\n"
                f"  seed={seed}\n"
                f"  expected={expected_sha}\n"
                f"  actual={actual_sha}"
            )

        threshold = load_yaml(
            path
        )

        if (
            threshold.get(
                "artifact_type"
            )
            != "resnet18_dev_derived_fpr10_threshold"
        ):

            raise RuntimeError(
                f"Unexpected threshold artifact type for seed {seed}."
            )

        if (
            threshold.get(
                "status"
            )
            != "DERIVED_UNDER_FROZEN_PROTOCOL"
        ):

            raise RuntimeError(
                f"Threshold artifact is not frozen for seed {seed}."
            )

        checkpoint = require_mapping(
            threshold[
                "checkpoint"
            ],
            f"threshold.{seed}.checkpoint",
        )

        if int(
            checkpoint[
                "seed"
            ]
        ) != seed:

            raise RuntimeError(
                f"Threshold/checkpoint seed mismatch for seed {seed}."
            )

        if (
            threshold[
                "resolution_selection"
            ][
                "selected_resolution"
            ]
            != "r512"
        ):

            raise RuntimeError(
                f"Threshold seed {seed} is not for r512."
            )

        boundary = require_mapping(
            threshold[
                "boundaries"
            ],
            f"threshold.{seed}.boundaries",
        )

        if boundary[
            "held_out_test_accessed"
        ] is not False:

            raise RuntimeError(
                f"Threshold seed {seed} reports held-out-test access."
            )

        fpr = require_mapping(
            threshold[
                "controlled_fpr10"
            ],
            f"threshold.{seed}.controlled_fpr10",
        )

        threshold_value = float(
            fpr[
                "threshold"
            ]
        )

        if not (
            threshold_value
            >= 0.0
        ):

            raise RuntimeError(
                f"Invalid frozen threshold for seed {seed}."
            )

        threshold_evidence[
            seed
        ] = {
            "path":
                relative_repo_path(
                    path
                ),

            "sha256":
                actual_sha,

            "threshold":
                threshold_value,

            "held_out_test_accessed":
                False,
        }

    return {
        "final_detection_protocol":
            {
                "path":
                    relative_repo_path(
                        protocol_path
                    ),

                "sha256":
                    actual_protocol_sha,
            },

        "dev_thresholds":
            threshold_evidence,
    }


# ======================================================================
# Workbook materialization
# ======================================================================

def read_test_metadata(
    *,
    workbook_path: Path,
    sheet_name: str,
    source_split: str,
    expected_total_rows: int,
) -> tuple[
    list[dict[str, Any]],
    Counter[str],
]:

    # read_only=True and data_only=True:
    # metadata cells only; no test JPEG/JSON path is followed.
    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:

        if sheet_name not in workbook.sheetnames:

            raise RuntimeError(
                f"Workbook missing sheet: {sheet_name}"
            )

        sheet = workbook[
            sheet_name
        ]

        rows = sheet.iter_rows(
            values_only=True
        )

        try:

            header_values = next(
                rows
            )

        except StopIteration as exc:

            raise RuntimeError(
                "Images sheet is empty."
            ) from exc

        headers = [
            normalized_text(
                value,
                label="Images header",
            )

            for value
            in header_values
        ]

        if len(
            headers
        ) != len(
            set(
                headers
            )
        ):

            raise RuntimeError(
                "Images sheet contains duplicate column names."
            )

        missing_columns = (
            REQUIRED_WORKBOOK_COLUMNS
            - set(
                headers
            )
        )

        if missing_columns:

            raise RuntimeError(
                "Images sheet missing required columns:\n"
                f"  {sorted(missing_columns)}"
            )

        index_by_name = {
            name:
                index

            for (
                index,
                name,
            ) in enumerate(
                headers
            )
        }

        all_row_count = 0

        split_counts: Counter[
            str
        ] = Counter()

        test_rows: list[
            dict[str, Any]
        ] = []

        seen_paths: set[
            str
        ] = set()

        seen_hashes: set[
            str
        ] = set()

        for workbook_row_number, values in enumerate(
            rows,
            start=2,
        ):

            # An entirely empty Excel row is ignored rather than counted
            # as a scientific image record.
            if all(
                value is None
                for value
                in values
            ):

                continue

            all_row_count += 1

            split = normalized_text(
                values[
                    index_by_name[
                        "split"
                    ]
                ],
                label=(
                    f"Images row {workbook_row_number} split"
                ),
            )

            split_counts[
                split
            ] += 1

            if split != source_split:

                continue

            traffic_type = normalized_text(
                values[
                    index_by_name[
                        "traffic_type"
                    ]
                ],
                label=(
                    f"Images row {workbook_row_number} traffic_type"
                ),
            )

            if traffic_type not in {
                "bonafide",
                "attack",
            }:

                raise RuntimeError(
                    "Unexpected test traffic_type:\n"
                    f"  workbook_row={workbook_row_number}\n"
                    f"  value={traffic_type!r}"
                )

            label = (
                0
                if traffic_type == "bonafide"
                else 1
            )

            variant = normalized_text(
                values[
                    index_by_name[
                        "variant"
                    ]
                ],
                label=(
                    f"Images row {workbook_row_number} variant"
                ),
                allow_empty=True,
            )

            image_path = normalized_text(
                values[
                    index_by_name[
                        "image_path"
                    ]
                ],
                label=(
                    f"Images row {workbook_row_number} image_path"
                ),
            )

            # Frozen workbook uses repository/dataset-relative POSIX paths.
            # Reject rather than silently repairing path syntax.
            if "\\" in image_path:

                raise RuntimeError(
                    "Test image_path contains backslash:\n"
                    f"  {image_path}"
                )

            pure_path = PurePosixPath(
                image_path
            )

            if pure_path.is_absolute():

                raise RuntimeError(
                    "Test image_path must be relative:\n"
                    f"  {image_path}"
                )

            if (
                not pure_path.parts
                or pure_path.parts[
                    0
                ] != "test"
            ):

                raise RuntimeError(
                    "Held-out manifest image_path must begin with test/:\n"
                    f"  {image_path}"
                )

            if any(
                part in {
                    ".",
                    "..",
                }
                for part
                in pure_path.parts
            ):

                raise RuntimeError(
                    "Unsafe test image_path:\n"
                    f"  {image_path}"
                )

            if image_path in seen_paths:

                raise RuntimeError(
                    "Duplicate held-out image_path:\n"
                    f"  {image_path}"
                )

            seen_paths.add(
                image_path
            )

            image_sha256 = normalized_text(
                values[
                    index_by_name[
                        "image_sha256"
                    ]
                ],
                label=(
                    f"Images row {workbook_row_number} image_sha256"
                ),
            ).lower()

            if SHA256_PATTERN.fullmatch(
                image_sha256
            ) is None:

                raise RuntimeError(
                    "Malformed test image SHA-256:\n"
                    f"  workbook_row={workbook_row_number}\n"
                    f"  value={image_sha256!r}"
                )

            if image_sha256 in seen_hashes:

                raise RuntimeError(
                    "Duplicate held-out image SHA-256:\n"
                    f"  {image_sha256}"
                )

            seen_hashes.add(
                image_sha256
            )

            test_rows.append(
                {
                    "source_workbook_row":
                        workbook_row_number,

                    "image_path":
                        image_path,

                    "image_sha256":
                        image_sha256,

                    "file_stem":
                        normalized_text(
                            values[
                                index_by_name[
                                    "file_stem"
                                ]
                            ],
                            label=(
                                f"Images row {workbook_row_number} "
                                "file_stem"
                            ),
                        ),

                    "traffic_type":
                        traffic_type,

                    "label":
                        label,

                    "variant":
                        variant,

                    "hardware_source":
                        normalized_text(
                            values[
                                index_by_name[
                                    "hardware_source"
                                ]
                            ],
                            label=(
                                f"Images row {workbook_row_number} "
                                "hardware_source"
                            ),
                        ),

                    "face_db":
                        normalized_text(
                            values[
                                index_by_name[
                                    "face_db"
                                ]
                            ],
                            label=(
                                f"Images row {workbook_row_number} "
                                "face_db"
                            ),
                        ),

                    "face_id":
                        normalized_text(
                            values[
                                index_by_name[
                                    "face_id"
                                ]
                            ],
                            label=(
                                f"Images row {workbook_row_number} "
                                "face_id"
                            ),
                        ),

                    "gender":
                        normalized_text(
                            values[
                                index_by_name[
                                    "gender"
                                ]
                            ],
                            label=(
                                f"Images row {workbook_row_number} "
                                "gender"
                            ),
                        ),

                    "project_role":
                        "held_out_test",
                }
            )

        if all_row_count != expected_total_rows:

            raise RuntimeError(
                "Frozen Images-sheet row-count mismatch:\n"
                f"  expected={expected_total_rows}\n"
                f"  actual={all_row_count}"
            )

        return (
            test_rows,
            split_counts,
        )

    finally:

        workbook.close()


# ======================================================================
# Exact count validation
# ======================================================================

def validate_counts(
    *,
    rows: list[dict[str, Any]],
    split_counts: Counter[str],
    config: Mapping[Any, Any],
) -> dict[str, Any]:

    source_cfg = require_mapping(
        config[
            "source_discovery"
        ],
        "source_discovery",
    )

    expected_split_counts = {
        str(
            key
        ):
            int(
                value
            )

        for (
            key,
            value,
        ) in require_mapping(
            source_cfg[
                "expected_split_counts"
            ],
            "source_discovery.expected_split_counts",
        ).items()
    }

    if dict(
        split_counts
    ) != expected_split_counts:

        raise RuntimeError(
            "Frozen Images-sheet split counts changed:\n"
            f"  expected={expected_split_counts}\n"
            f"  actual={dict(split_counts)}"
        )

    test_cfg = require_mapping(
        config[
            "held_out_test"
        ],
        "held_out_test",
    )

    expected_rows = int(
        test_cfg[
            "expected_rows"
        ]
    )

    if len(
        rows
    ) != expected_rows:

        raise RuntimeError(
            "Held-out-test row count mismatch:\n"
            f"  expected={expected_rows}\n"
            f"  actual={len(rows)}"
        )

    observed_class_counts = Counter(
        str(
            row[
                "traffic_type"
            ]
        )

        for row
        in rows
    )

    expected_class_counts = {
        str(
            key
        ):
            int(
                value
            )

        for (
            key,
            value,
        ) in require_mapping(
            test_cfg[
                "expected_class_counts"
            ],
            "held_out_test.expected_class_counts",
        ).items()
    }

    if dict(
        observed_class_counts
    ) != expected_class_counts:

        raise RuntimeError(
            "Held-out class-count mismatch:\n"
            f"  expected={expected_class_counts}\n"
            f"  actual={dict(observed_class_counts)}"
        )

    observed_groups = Counter(
        (
            str(
                row[
                    "traffic_type"
                ]
            ),
            str(
                row[
                    "variant"
                ]
            ),
        )

        for row
        in rows
    )

    expected_groups_cfg = require_mapping(
        test_cfg[
            "expected_groups"
        ],
        "held_out_test.expected_groups",
    )

    expected_groups: dict[
        tuple[str, str],
        int,
    ] = {}

    for traffic_type, variant_map in (
        expected_groups_cfg.items()
    ):

        variant_mapping = require_mapping(
            variant_map,
            (
                "held_out_test.expected_groups."
                f"{traffic_type}"
            ),
        )

        for variant, count in (
            variant_mapping.items()
        ):

            expected_groups[
                (
                    str(
                        traffic_type
                    ),
                    str(
                        variant
                    ),
                )
            ] = int(
                count
            )

    if dict(
        observed_groups
    ) != expected_groups:

        raise RuntimeError(
            "Held-out traffic/variant composition mismatch:\n"
            f"  expected={expected_groups}\n"
            f"  actual={dict(observed_groups)}"
        )

    # Explicitly reconcile the frozen numeric encoding.
    class_mapping = {
        str(
            key
        ):
            int(
                value
            )

        for (
            key,
            value,
        ) in require_mapping(
            test_cfg[
                "class_mapping"
            ],
            "held_out_test.class_mapping",
        ).items()
    }

    if class_mapping != {
        "bonafide":
            0,

        "attack":
            1,
    }:

        raise RuntimeError(
            "Frozen T2 class mapping changed."
        )

    for row in rows:

        expected_label = class_mapping[
            row[
                "traffic_type"
            ]
        ]

        if int(
            row[
                "label"
            ]
        ) != expected_label:

            raise RuntimeError(
                "Semantic/numerical ground-truth mismatch."
            )

    hardware_counts = Counter(
        str(
            row[
                "hardware_source"
            ]
        )

        for row
        in rows
    )

    return {
        "class_counts":
            dict(
                observed_class_counts
            ),

        "traffic_variant_counts":
            {
                f"{traffic_type}|{variant}":
                    count

                for (
                    (
                        traffic_type,
                        variant,
                    ),
                    count,
                ) in sorted(
                    observed_groups.items()
                )
            },

        "hardware_counts":
            dict(
                sorted(
                    hardware_counts.items()
                )
            ),

        "unique_image_paths":
            len(
                {
                    row[
                        "image_path"
                    ]
                    for row
                    in rows
                }
            ),

        "unique_image_sha256":
            len(
                {
                    row[
                        "image_sha256"
                    ]
                    for row
                    in rows
                }
            ),

        "unique_file_stems":
            len(
                {
                    row[
                        "file_stem"
                    ]
                    for row
                    in rows
                }
            ),
    }


# ======================================================================
# Artifact writing
# ======================================================================

def write_csv(
    *,
    path: Path,
    rows: list[dict[str, Any]],
) -> None:

    with path.open(
        "x",
        encoding="utf-8",
        newline="",
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=OUTPUT_COLUMNS,
            extrasaction="raise",
            lineterminator="\n",
        )

        writer.writeheader()

        writer.writerows(
            rows
        )


# ======================================================================
# Main
# ======================================================================

def main() -> int:

    parser = argparse.ArgumentParser(
        description=(
            "Materialize canonical FantasyID held-out-test metadata "
            "manifest without opening test images."
        )
    )

    parser.add_argument(
        "--config",
        default=(
            "tools/"
            "finalize_fantasyid_test_manifest_config.yaml"
        ),
    )

    args = parser.parse_args()

    git_commit, git_branch = (
        require_clean_git()
    )

    config_path = resolve_repo_path(
        args.config
    )

    config = load_yaml(
        config_path
    )

    # Validate pre-test decisions before creating this run's own log.
    prerequisite_evidence = (
        validate_prerequisites(
            config=config
        )
    )

    source_cfg = require_mapping(
        config[
            "source_discovery"
        ],
        "source_discovery",
    )

    workbook_path = resolve_repo_path(
        source_cfg[
            "workbook"
        ]
    )

    require_git_tracked(
        workbook_path
    )

    expected_workbook_sha = str(
        source_cfg[
            "sha256"
        ]
    )

    actual_workbook_sha = sha256_file(
        workbook_path
    )

    if actual_workbook_sha != expected_workbook_sha:

        raise RuntimeError(
            "Frozen discovery workbook SHA mismatch:\n"
            f"  expected={expected_workbook_sha}\n"
            f"  actual={actual_workbook_sha}"
        )

    log_path = configure_logging(
        config=config
    )

    output_cfg = require_mapping(
        config[
            "output"
        ],
        "output",
    )

    output_directory = resolve_repo_path(
        output_cfg[
            "directory"
        ]
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    csv_path = (
        output_directory
        / str(
            output_cfg[
                "csv_filename"
            ]
        )
    )

    metadata_path = (
        output_directory
        / str(
            output_cfg[
                "metadata_filename"
            ]
        )
    )

    csv_partial = Path(
        str(
            csv_path
        )
        + ".partial"
    )

    metadata_partial = Path(
        str(
            metadata_path
        )
        + ".partial"
    )

    try:

        LOGGER.info(
            "=" * 72
        )

        LOGGER.info(
            "FANTASYID HELD-OUT TEST MANIFEST FINALIZATION"
        )

        LOGGER.info(
            "=" * 72
        )

        LOGGER.info(
            "Git commit: %s",
            git_commit,
        )

        LOGGER.info(
            "Git branch: %s",
            git_branch,
        )

        LOGGER.info(
            "Discovery workbook: %s",
            workbook_path,
        )

        LOGGER.info(
            "Discovery workbook SHA-256: %s",
            actual_workbook_sha,
        )

        if (
            csv_path.exists()
            or metadata_path.exists()
            or csv_partial.exists()
            or metadata_partial.exists()
        ):

            raise FileExistsError(
                "Canonical test-manifest output already exists."
            )

        for seed in (
            8,
            9,
            10,
        ):

            item = prerequisite_evidence[
                "dev_thresholds"
            ][
                seed
            ]

            LOGGER.info(
                "[PASS] seed %d dev threshold frozen before "
                "test-manifest materialization | threshold=%.17g | "
                "sha256=%s",
                seed,
                item[
                    "threshold"
                ],
                item[
                    "sha256"
                ],
            )

        test_cfg = require_mapping(
            config[
                "held_out_test"
            ],
            "held_out_test",
        )

        rows, split_counts = (
            read_test_metadata(
                workbook_path=(
                    workbook_path
                ),
                sheet_name=str(
                    source_cfg[
                        "sheet"
                    ]
                ),
                source_split=str(
                    test_cfg[
                        "source_split"
                    ]
                ),
                expected_total_rows=int(
                    source_cfg[
                        "expected_total_rows"
                    ]
                ),
            )
        )

        composition = validate_counts(
            rows=rows,
            split_counts=split_counts,
            config=config,
        )

        LOGGER.info(
            "[PASS] frozen Images sheet rows = 3284"
        )

        LOGGER.info(
            "[PASS] frozen split counts = train 1899 / test 1385"
        )

        LOGGER.info(
            "[PASS] held-out rows = 1385"
        )

        LOGGER.info(
            "[PASS] ground truth = bonafide 300 / attack 1085"
        )

        LOGGER.info(
            "[PASS] T2 label encoding = bonafide 0 / attack 1"
        )

        LOGGER.info(
            "[PASS] test group = bonafide 300"
        )

        LOGGER.info(
            "[PASS] test attack digital_3 = 786"
        )

        LOGGER.info(
            "[PASS] test attack facedancer = 150"
        )

        LOGGER.info(
            "[PASS] test attack textdiffuserft_bfei = 149"
        )

        LOGGER.info(
            "[PASS] unique image paths = %d",
            composition[
                "unique_image_paths"
            ],
        )

        LOGGER.info(
            "[PASS] unique frozen image SHA-256 values = %d",
            composition[
                "unique_image_sha256"
            ],
        )

        LOGGER.info(
            "Observed hardware counts: %s",
            composition[
                "hardware_counts"
            ],
        )

        LOGGER.info(
            "Observed unique file stems: %d",
            composition[
                "unique_file_stems"
            ],
        )

        # --------------------------------------------------------------
        # Write canonical CSV in frozen workbook order.
        # --------------------------------------------------------------

        write_csv(
            path=csv_partial,
            rows=rows,
        )

        csv_sha = sha256_file(
            csv_partial
        )

        metadata = {
            "schema_version":
                1,

            "artifact_type":
                "fantasyid_held_out_test_manifest",

            "status":
                "FROZEN_BEFORE_MODEL_TEST_ACCESS",

            "created_at_utc":
                datetime.now(
                    timezone.utc
                ).isoformat(),

            "git":
                {
                    "commit_sha":
                        git_commit,

                    "branch":
                        git_branch,
                },

            "source_discovery":
                {
                    "workbook":
                        relative_repo_path(
                            workbook_path
                        ),

                    "sha256":
                        actual_workbook_sha,

                    "sheet":
                        str(
                            source_cfg[
                                "sheet"
                            ]
                        ),

                    "source_total_rows":
                        int(
                            source_cfg[
                                "expected_total_rows"
                            ]
                        ),

                    "source_split_counts":
                        dict(
                            split_counts
                        ),
                },

            "held_out_test":
                {
                    "source_split":
                        "test",

                    "project_role":
                        "held_out_test",

                    "rows":
                        len(
                            rows
                        ),

                    "class_mapping":
                        {
                            "bonafide":
                                0,

                            "attack":
                                1,
                        },

                    "class_counts":
                        composition[
                            "class_counts"
                        ],

                    "traffic_variant_counts":
                        composition[
                            "traffic_variant_counts"
                        ],

                    "hardware_counts_observed_metadata_only":
                        composition[
                            "hardware_counts"
                        ],

                    "unique_file_stems":
                        composition[
                            "unique_file_stems"
                        ],

                    "unique_image_paths":
                        composition[
                            "unique_image_paths"
                        ],

                    "unique_frozen_image_sha256":
                        composition[
                            "unique_image_sha256"
                        ],
                },

            "manifest":
                {
                    "path":
                        relative_repo_path(
                            csv_path
                        ),

                    "sha256":
                        csv_sha,

                    "rows":
                        len(
                            rows
                        ),

                    "columns":
                        list(
                            OUTPUT_COLUMNS
                        ),

                    "order":
                        (
                            "frozen_discovery_Images_sheet_order_"
                            "filtered_to_split_test"
                        ),

                    "sorting_performed":
                        False,
                },

            "pre_test_prerequisites":
                prerequisite_evidence,

            "access_boundary":
                {
                    "test_workbook_metadata_read":
                        True,

                    "test_image_paths_followed":
                        False,

                    "test_image_files_opened":
                        False,

                    "test_image_pixels_decoded":
                        False,

                    "test_image_hashes_recomputed":
                        False,

                    "test_json_files_opened":
                        False,

                    "test_regions_inspected":
                        False,

                    "model_loaded":
                        False,

                    "checkpoint_loaded":
                        False,

                    "model_inference_performed":
                        False,

                    "metric_calculation_performed":
                        False,

                    "threshold_derived_or_modified":
                        False,
                },
        }

        with metadata_partial.open(
            "x",
            encoding="utf-8",
        ) as file:

            yaml.safe_dump(
                metadata,
                file,
                sort_keys=False,
            )

        metadata_sha = sha256_file(
            metadata_partial
        )

        # Promote only after both artifacts are complete.
        promoted_csv = False

        try:

            os.replace(
                csv_partial,
                csv_path,
            )

            promoted_csv = True

            os.replace(
                metadata_partial,
                metadata_path,
            )

        except Exception:

            if (
                promoted_csv
                and csv_path.exists()
            ):

                csv_path.unlink()

            raise

        LOGGER.info(
            ""
        )

        LOGGER.info(
            "[PASS] canonical held-out manifest: %s",
            csv_path,
        )

        LOGGER.info(
            "Manifest SHA-256: %s",
            csv_sha,
        )

        LOGGER.info(
            "[PASS] manifest metadata: %s",
            metadata_path,
        )

        LOGGER.info(
            "Metadata SHA-256: %s",
            metadata_sha,
        )

        LOGGER.info(
            ""
        )

        LOGGER.info(
            "[PASS] test image files opened: FALSE"
        )

        LOGGER.info(
            "[PASS] test pixels decoded: FALSE"
        )

        LOGGER.info(
            "[PASS] test JSON/regions inspected: FALSE"
        )

        LOGGER.info(
            "[PASS] model/checkpoint loaded: FALSE"
        )

        LOGGER.info(
            "[PASS] model inference performed: FALSE"
        )

        LOGGER.info(
            "[PASS] threshold derivation/modification: FALSE"
        )

        LOGGER.info(
            "=" * 72
        )

        LOGGER.info(
            "HELD-OUT TEST MANIFEST: FROZEN BEFORE MODEL TEST ACCESS"
        )

        LOGGER.info(
            "=" * 72
        )

        return 0

    except Exception:

        LOGGER.exception(
            "HELD-OUT TEST MANIFEST FINALIZATION: FAIL"
        )

        for partial in (
            csv_partial,
            metadata_partial,
        ):

            if partial.exists():

                partial.unlink()

        return 1

    finally:

        for handler in list(
            LOGGER.handlers
        ):

            handler.flush()
            handler.close()

        LOGGER.handlers.clear()


if __name__ == "__main__":

    raise SystemExit(
        main()
    )